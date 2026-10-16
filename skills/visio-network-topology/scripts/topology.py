"""Portable network-topology planning, Visio rendering and verification (v0.1).

Planning uses only Python plus openpyxl for XLSX. Rendering imports pywin32 lazily.
User inputs, stencils and generated artifacts are never uploaded by this module.
"""
from __future__ import annotations

import argparse
import collections
import csv
import json
import math
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

VERSION = 'v0.1'
NS = {'v': 'http://schemas.microsoft.com/office/visio/2012/main'}
HEADERS = ['本端设备名', '本端设备类型', '本端接口', '对端设备名', '对端设备类型', '对端接口', '线路类型']
FIELDS = ['from', 'from_type', 'from_port', 'to', 'to_type', 'to_port', 'line_type']
DEFAULT_TYPES = {
    'core-switch': ('core', 'Core-Switch'), 'core switch': ('core', 'Core-Switch'),
    '核心交换机': ('core', 'Core-Switch'), 'switch': ('access', 'Switch'),
    '交换机': ('access', 'Switch'), '接入交换机': ('access', 'Switch'),
    'firewall': ('firewall', 'Firewall'), '防火墙': ('firewall', 'Firewall'),
}
COLORS = {'40G': '#FF0000', '10G': '#FFD000', 'copper': '#000000'}
LEGEND = [('40G', '40G 光纤'), ('10G', '10G 光纤'), ('copper', '网线')]


class TopologyError(ValueError):
    """An input or layout cannot be rendered faithfully without adjustment."""


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def natural(value):
    return [int(s) if s.isdigit() else s.lower() for s in re.split(r'(\d+)', value)]


def read_rows(path, sheet=None):
    """Read exact headers; source cells are data, never executable instructions."""
    path = Path(path)
    if path.suffix.lower() == '.json':
        value = read_json(path)
        rows = value.get('rows', value.get('links', [])) if isinstance(value, dict) else value
        return rows, {'file': path.name, 'sheet': None}
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            table = list(csv.reader(stream))
        sheet_name = None
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet and sheet not in workbook.sheetnames:
                raise TopologyError(f'Worksheet not found: {sheet}')
            candidates = [workbook[sheet]] if sheet else list(workbook.worksheets)
            chosen = next((s for s in candidates if any(c.value is not None for row in s.iter_rows() for c in row)), None)
            if chosen is None:
                raise TopologyError('Workbook contains no data')
            sheet_name = chosen.title
            # Reject formulas rather than silently consuming missing/stale cached results.
            formulas = load_workbook(path, read_only=True, data_only=False)
            try:
                if any(c.data_type == 'f' for row in formulas[sheet_name].iter_rows() for c in row):
                    raise TopologyError('Formula cells found; export the link table as values before planning')
            finally:
                formulas.close()
            table = list(chosen.iter_rows(values_only=True))
        finally:
            workbook.close()
    else:
        raise TopologyError('Input must be .xlsx, .csv or .json')
    header_index = next((i for i, row in enumerate(table) if any(v is not None and str(v).strip() for v in row)), None)
    if header_index is None:
        raise TopologyError('Input table is empty')
    headers = [str(v).strip() if v is not None else '' for v in table[header_index]]
    missing = [h for h in HEADERS if h not in headers]
    if missing:
        raise TopologyError(f'Missing column(s): {missing}')
    if any(headers.count(h) != 1 for h in HEADERS):
        raise TopologyError('Required column names must be unique')
    indices = [headers.index(h) for h in HEADERS]
    rows = []
    for number, row in enumerate(table[header_index + 1:], header_index + 2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        values = [row[i] if i < len(row) else None for i in indices]
        rows.append(dict(zip(FIELDS, values), source_row=number))
    return rows, {'file': path.name, 'sheet': sheet_name}


def line_class(value):
    compact = re.sub(r'\s+', '', str(value)).lower()
    if compact in ('40g', '40g光纤', '40gbps'):
        return '40G'
    if compact in ('10g', '10g光纤', '10gbps'):
        return '10G'
    if compact in ('网线', 'copper', 'ethernet', '铜缆'):
        return 'copper'
    raise TopologyError(f'Unsupported line type: {value!r}; configure line_types explicitly')


def normalize(rows, config):
    if not isinstance(rows, list) or not rows:
        raise TopologyError('No link records found')
    devices, links, seen, port_peers, duplicate_rows = {}, [], {}, {}, []
    types = dict(DEFAULT_TYPES)
    for kind, definition in config.get('device_types', {}).items():
        role = definition['role']
        if role not in ('firewall', 'core', 'access'):
            raise TopologyError(f'Unsupported role: {role}')
        types[kind.strip().lower()] = (role, definition['master'])
    for number, raw in enumerate(rows, 2):
        row = {k: str(raw.get(k, '')).strip() if raw.get(k) is not None else '' for k in FIELDS}
        source_row = raw.get('source_row', raw.get('excel_row', number))
        if any(not row[k] for k in FIELDS):
            raise TopologyError(f'Row {source_row}: all seven fields are required')
        if row['from'] == row['to']:
            raise TopologyError(f'Row {source_row}: self-links need a manual layout')
        for side in ('from', 'to'):
            name, kind = row[side], row[side + '_type'].lower()
            if kind not in types:
                raise TopologyError(f'Row {source_row}: unknown device type {kind!r}')
            role, master = types[kind]
            descriptor = {'name': name, 'role': role, 'master': master}
            if name in devices and devices[name] != descriptor:
                raise TopologyError(f'Conflicting types for device {name!r}')
            devices[name] = descriptor
        configured = config.get('line_types', {}).get(row['line_type'])
        category = configured if configured is not None else line_class(row['line_type'])
        if category not in COLORS:
            raise TopologyError(f'Line mapping must resolve to one of {list(COLORS)}')
        endpoints = sorted([(row['from'], row['from_port']), (row['to'], row['to_port'])])
        key = tuple(endpoints)
        if key in seen:
            existing = links[seen[key]]
            if existing['category'] != category:
                raise TopologyError(f'Conflicting line types for {key}')
            existing['source_rows'].append(source_row)
            duplicate_rows.append(source_row)
            continue
        for index, endpoint in enumerate(endpoints):
            peer = endpoints[1-index]
            if endpoint in port_peers and port_peers[endpoint] != peer:
                raise TopologyError(f'Physical port reused with another peer: {endpoint}')
            port_peers[endpoint] = peer
        seen[key] = len(links)
        links.append(dict(row, id=f'L{len(links)+1:04d}', category=category,
                          color=COLORS[category], source_rows=[source_row]))
    ordered = sorted(devices.values(), key=lambda d: (('firewall', 'core', 'access').index(d['role']), natural(d['name'])))
    for index, device in enumerate(ordered, 1):
        device['id'] = f'D{index:04d}'
    return ordered, links, duplicate_rows


def rect(node, padding=0):
    return (node['x']-padding, node['y']-padding, node['x']+node['w']+padding, node['y']+node['h']+padding)


def overlap(a, b, padding=.025):
    return a[2]+padding>b[0] and b[2]+padding>a[0] and a[3]+padding>b[1] and b[3]+padding>a[1]


def segment_hits(p, q, box):
    dx, dy = q[0]-p[0], q[1]-p[1]
    low, high = 0.0, 1.0
    for coefficient, distance in ((-dx,p[0]-box[0]), (dx,box[2]-p[0]), (-dy,p[1]-box[1]), (dy,box[3]-p[1])):
        if abs(coefficient)<1e-12:
            if distance<0:
                return False
        else:
            ratio=distance/coefficient
            if coefficient<0:
                low=max(low,ratio)
            else:
                high=min(high,ratio)
            if low>high:
                return False
    return True


def name_box(node, width):
    x, y, w = node['cx']-1, node['y']-.4, 2.0
    if node['role']=='firewall':
        x, y, w = node['cx']-.8, node['y']+node['h']+.14, 1.6
    elif node['role']=='core':
        x = node['x']-2.2 if node['cx']<width/2 else node['x']+node['w']+.2
        y = node['cy']-.16
    return x,y,x+w,y+.32


def anchor(node, value):
    return node['x']+node['w']*value[0], node['y']+node['h']*value[1]


def master_names(path):
    with ZipFile(path) as archive:
        root=ET.fromstring(archive.read('visio/masters/masters.xml'))
    return [m.attrib.get('Name', m.attrib.get('NameU')) for m in root.findall('v:Master',NS)]


def plan(rows, source, config=None, stencil=None):
    config = config or {}
    nodes, links, duplicates = normalize(rows,config)
    lookup={n['name']:n for n in nodes}
    groups={role:[n for n in nodes if n['role']==role] for role in ('firewall','core','access')}
    width=max(18.0,len(groups['access'])*3.2+4,len(groups['core'])*5.0+4,len(groups['firewall'])*3.4+4)
    width=float(config.get('page_width',width))
    height=float(config.get('page_height',21))
    if width<10 or height<16:
        raise TopologyError('Canvas is too small; enlarge the page rather than shrinking text')
    degree=collections.Counter(name for link in links for name in (link['from'],link['to']))
    levels={'firewall':height-4.2,'core':height-9.1,'access':3.9}
    for role, group in groups.items():
        for index,node in enumerate(group):
            size={'firewall':(1.6,.6025),'access':(1.25,1.02429),'core':(1.6,1.31115)}[role]
            if role=='core' and degree[node['name']]>=8:
                size=(3.4,3.4*1.31115/1.6)
            if role=='core' and len(group)==2:
                cx=width*(11/36 if index==0 else 25/36)
            else:
                spacing=3.4 if role=='firewall' else (3.2 if role=='access' else 5.0)
                # Paired firewalls remain compact; different pairs get additional space.
                slots=[i*spacing+(i//2)*1.0 for i in range(len(group))] if role=='firewall' else [i*spacing for i in range(len(group))]
                cx=width/2+slots[index]-(slots[0]+slots[-1])/2
            placement=config.get('positions',{}).get(node['name'],{})
            w,h=float(placement.get('w',size[0])),float(placement.get('h',size[1]))
            cx,cy=float(placement.get('cx',cx)),float(placement.get('cy',levels[role]))
            if w<=0 or h<=0:
                raise TopologyError(f'Invalid dimensions for {node["name"]}')
            node.update(cx=cx,cy=cy,w=w,h=h,x=cx-w/2,y=cy-h/2)
            node['name_box']=name_box(node,width)
            if min(node['x'],node['y'])<.1 or node['x']+w>width-.1 or node['y']+h>height-.1:
                raise TopologyError(f'Device outside canvas: {node["name"]}')
    for i,node in enumerate(nodes):
        for other in nodes[i+1:]:
            if overlap(rect(node,.12),rect(other,.12),0):
                raise TopologyError(f'Devices overlap: {node["name"]}, {other["name"]}')
    edge_ports=collections.defaultdict(list)
    same_level=collections.defaultdict(list)
    for link in links:
        a,b=lookup[link['from']],lookup[link['to']]
        if abs(a['cy']-b['cy'])<1e-6:
            same_level[tuple(sorted((a['name'],b['name'])))].append(link)
        else:
            for side,other in (('from',b),('to',a)):
                current=lookup[link[side]]
                edge='top' if other['cy']>current['cy'] else 'bottom'
                edge_ports[(current['name'],edge)].append((other['cx'],link['id'],side,link))
    for (name,edge),entries in edge_ports.items():
        entries.sort(key=lambda item:(item[0],item[1]))
        for i,(_,_,side,link) in enumerate(entries):
            fraction=.5 if len(entries)==1 else (.3+.4*i/(len(entries)-1) if len(entries)<=2 else .06+.88*i/(len(entries)-1))
            link[side+'_anchor']=[fraction,1 if edge=='top' else 0]
    for entries in same_level.values():
        for i,link in enumerate(entries):
            a,b=lookup[link['from']],lookup[link['to']]
            fraction=.5 if len(entries)==1 else .72-.44*i/(len(entries)-1)
            link['from_anchor']=[1 if a['cx']<b['cx'] else 0,fraction]
            link['to_anchor']=[0 if a['cx']<b['cx'] else 1,fraction]
    for link in links:
        a,b=lookup[link['from']],lookup[link['to']]
        link['p'],link['q']=anchor(a,link['from_anchor']),anchor(b,link['to_anchor'])
        for node in nodes:
            if node['name'] not in (a['name'],b['name']) and segment_hits(link['p'],link['q'],rect(node)):
                raise TopologyError(f'{link["id"]} passes through {node["name"]}; adjust positions in config')
    labels=place_labels(nodes,links,width)
    stencil_path=str(Path(stencil).resolve()) if stencil else None
    if stencil_path:
        masters={name.lower():name for name in master_names(stencil_path)}
        for node in nodes:
            if node['master'].lower() not in masters:
                raise TopologyError(f'Master {node["master"]!r} not in stencil; available: {list(masters.values())}')
            node['master']=masters[node['master'].lower()]
    result={'version':VERSION,'title':config.get('title','网络拓扑图'),'source':source,
            'page_width':width,'page_height':height,'stencil':stencil_path,'nodes':nodes,'links':links,
            'labels':labels,'legend':[{'category':c,'text':t,'color':COLORS[c]} for c,t in LEGEND],
            'duplicate_source_rows':duplicates}
    return result


def place_labels(nodes,links,page_width):
    lookup={n['name']:n for n in nodes}
    obstacles=[rect(n,.035) for n in nodes]+[n['name_box'] for n in nodes]
    degree=collections.Counter(name for link in links for name in (link['from'],link['to']))
    requests=[]
    for link in links:
        p,q=link['p'],link['q']
        dx,dy=q[0]-p[0],q[1]-p[1]
        length=math.hypot(dx,dy)
        if length<.8:
            raise TopologyError(f'{link["id"]}: line too short for two endpoint labels')
        theta=math.atan2(dy,dx)
        for side in ('from','to'):
            requests.append((link,side,dx,dy,length,theta))
    requests.sort(key=lambda r:(r[0]['category']=='10G',-degree[r[0][r[1]]],r[0][r[1]],r[5]))
    labels,boxes=[],[]
    for link,side,dx,dy,length,theta in requests:
        ux,uy=dx/length,dy/length
        width,height=.65,.23
        ex,ey=abs(ux)*width/2+abs(uy)*height/2,abs(uy)*width/2+abs(ux)*height/2
        placed=False
        for distance in [.43+.12*i for i in range(44)]:
            if distance>length*.45:
                continue
            fraction=distance/length if side=='from' else 1-distance/length
            cx,cy=link['p'][0]+dx*fraction,link['p'][1]+dy*fraction
            box=(cx-ex,cy-ey,cx+ex,cy+ey)
            if any(overlap(box,b) for b in obstacles+boxes):
                continue
            crossed=False
            for other in links:
                if other['id']==link['id']:
                    continue
                def local(point):
                    x,y=point[0]-cx,point[1]-cy
                    return x*ux+y*uy,-x*uy+y*ux
                if segment_hits(local(other['p']),local(other['q']),(-width/2-.045,-height/2-.045,width/2+.045,height/2+.045)):
                    crossed=True
                    break
            if crossed:
                continue
            labels.append({'id':link['id']+'_'+side+'_port','link':link['id'],'side':side,'device':link[side],
                           'text':link[side+'_port'],'fraction':fraction,'x':cx,'y':cy,'theta':theta,
                           'width':width,'height':height,'endpoint_distance':distance,'color':link['color']})
            boxes.append(box)
            placed=True
            break
        if not placed:
            raise TopologyError(f'No clear inline label position for {link["id"]} {side}; enlarge canvas or spread device positions')
    return labels


def validate_plan(model):
    nodes,links,labels=model['nodes'],model['links'],model['labels']
    if len(labels)!=2*len(links):
        raise TopologyError('Each physical link must have exactly two port labels')
    ids=[item['id'] for item in nodes+links+labels]
    if len(set(ids))!=len(ids):
        raise TopologyError('Shape IDs must be unique')
    by_link={item['id']:item for item in links}
    boxes=[]
    for tag in labels:
        link=by_link[tag['link']]
        p,q=link['p'],link['q']
        dx,dy=q[0]-p[0],q[1]-p[1]
        length=math.hypot(dx,dy)
        ux,uy=dx/length,dy/length
        distance=abs(dx*(tag['y']-p[1])-dy*(tag['x']-p[0]))/length
        if distance>1e-7 or tag['text']!=link[tag['side']+'_port']:
            raise TopologyError(f'Incorrect inline port label: {tag["id"]}')
        ex=abs(ux)*tag['width']/2+abs(uy)*tag['height']/2
        ey=abs(uy)*tag['width']/2+abs(ux)*tag['height']/2
        box=(tag['x']-ex,tag['y']-ey,tag['x']+ex,tag['y']+ey)
        if any(overlap(box,b,padding=0) for b in boxes):
            raise TopologyError(f'Overlapping port labels near {tag["id"]}')
        if any(overlap(box,rect(n),0) or overlap(box,n['name_box'],0) for n in nodes):
            raise TopologyError(f'Port label obscures device/name: {tag["id"]}')
        for other in links:
            if other['id']==link['id']:
                continue
            def local(point):
                x,y=point[0]-tag['x'],point[1]-tag['y']
                return x*ux+y*uy,-x*uy+y*ux
            if segment_hits(local(other['p']),local(other['q']),(-tag['width']/2,-tag['height']/2,tag['width']/2,tag['height']/2)):
                raise TopologyError(f'Another line crosses {tag["id"]}')
        boxes.append(box)
    lookup={n['name']:n for n in nodes}
    for index,node in enumerate(nodes):
        if any(overlap(rect(node),rect(other),0) for other in nodes[index+1:]):
            raise TopologyError('Device bodies overlap')
    for link in links:
        if link['color']!=COLORS[link['category']]:
            raise TopologyError(f'Incorrect line color: {link["id"]}')
        for node in nodes:
            if node['name'] not in (link['from'],link['to']) and segment_hits(link['p'],link['q'],rect(node)):
                raise TopologyError(f'Line crosses an unrelated device: {link["id"]}')
    return {'version':VERSION,'devices':len(nodes),'links':len(links),'port_labels':len(labels),
            'line_types':dict(collections.Counter(l['category'] for l in links)),
            'duplicate_source_rows':model.get('duplicate_source_rows',[]),
            'overlapping_devices':0,'overlapping_port_labels':0,'foreign_lines_crossing_tags':0,
            'max_port_distance_inches':round(max(t['endpoint_distance'] for t in labels),2)}


def verify_vsdx(path,model):
    report=validate_plan(model)
    with ZipFile(path) as archive:
        page=ET.fromstring(archive.read('visio/pages/page1.xml'))
    shapes={s.attrib.get('NameU'):s for s in page.findall('v:Shapes/v:Shape',NS)}
    for node in model['nodes']:
        if node['id'] not in shapes:
            raise TopologyError(f'Missing saved device: {node["name"]}')
        cells={c.attrib['N']:c.attrib.get('V') for c in shapes[node['id']].findall('v:Cell',NS)}
        for cell,key in [('PinX','cx'),('PinY','cy'),('Width','w'),('Height','h')]:
            if abs(float(cells[cell])-node[key])>1e-6:
                raise TopologyError(f'Saved device position/size differs: {node["name"]}')
    for link in model['links']:
        shape=shapes[link['id']]
        geometry=[r.attrib.get('T') for sec in shape.findall('v:Section',NS) if sec.attrib.get('N')=='Geometry' for r in sec.findall('v:Row',NS)]
        if geometry!=['MoveTo','LineTo']:
            raise TopologyError(f'{link["id"]} is not a single straight segment')
        cells={c.attrib['N']:c.attrib.get('V') for c in shape.findall('v:Cell',NS)}
        if cells.get('LineColor','').upper()!=link['color']:
            raise TopologyError(f'Saved color mismatch: {link["id"]}')
    for tag in model['labels']:
        shape=shapes[tag['id']]
        value=''.join(shape.find('v:Text',NS).itertext()).strip()
        cells={c.attrib['N']:c.attrib.get('V') for c in shape.findall('v:Cell',NS)}
        if value!=tag['text'] or abs(float(cells['PinX'])-tag['x'])>1e-6 or abs(float(cells['PinY'])-tag['y'])>1e-6:
            raise TopologyError(f'Saved port label mismatch: {tag["id"]}')
    glue=page.findall('v:Connects/v:Connect',NS)
    if len(glue)!=2*len(model['links']):
        raise TopologyError('Saved link endpoints are not all glued to devices')
    legend=[]
    for i,item in enumerate(model['legend']):
        shape=shapes[f'LegendText_{i}']
        if ''.join(shape.find('v:Text',NS).itertext()).strip()!=item['text']:
            raise TopologyError('Legend text/order mismatch')
        cells={c.attrib['N']:c.attrib.get('V') for c in shape.findall('v:Cell',NS)}
        legend.append((float(cells['PinX']),float(cells['PinY'])))
    if not (max(x for x,y in legend)<5 and legend[0][1]>legend[1][1]>legend[2][1]):
        raise TopologyError('Legend must be vertical at the upper left')
    report.update(saved_straight_segments=len(model['links']),glued_endpoints=len(glue),saved_port_labels=len(model['labels']),
                  legend_order=[v['text'] for v in model['legend']])
    return report


def rgb(color):
    value=color.lstrip('#')
    if not re.fullmatch(r'[0-9A-Fa-f]{6}',value):
        raise TopologyError(f'Invalid color: {color}')
    return 'RGB('+','.join(str(int(value[i:i+2],16)) for i in (0,2,4))+')'


def render(model,output,preview=None,visible=False):
    validate_plan(model)
    if sys.platform!='win32':
        raise TopologyError('Rendering requires Windows and installed Microsoft Visio; plan/validate remain portable')
    import pythoncom
    import win32com.client
    output=Path(output).resolve()
    output.parent.mkdir(parents=True,exist_ok=True)
    if preview:
        Path(preview).resolve().parent.mkdir(parents=True,exist_ok=True)
    pythoncom.CoInitialize()
    app=doc=stencil=None
    try:
        print('Connecting to Microsoft Visio...',flush=True)
        app=win32com.client.DispatchEx('Visio.Application')
        initially_empty=app.Documents.Count==0
        if initially_empty:
            app.Visible=-1 if visible else 0
        doc=app.Documents.Add('')
        page=doc.Pages.Item(1)
        page.Name='网络拓扑'
        for cell,key in [('PageWidth','page_width'),('PageHeight','page_height')]:
            page.PageSheet.CellsU(cell).ResultIU=model[key]
        if model.get('stencil'):
            stencil=app.Documents.OpenEx(str(Path(model['stencil']).resolve()),66)
        font_id=None
        for family in ('Microsoft YaHei','微软雅黑','Arial'):
            try:
                font_id=doc.Fonts.Item(family).ID
                break
            except Exception:
                pass
        def label(value,box,size=12,bold=False):
            shape=page.DrawRectangle(*box)
            shape.Text=value
            shape.CellsU('LinePattern').ResultIU=0
            shape.CellsU('FillPattern').ResultIU=0
            shape.CellsU('Char.Size').FormulaU=f'{size} pt'
            if font_id is not None:
                shape.CellsU('Char.Font').ResultIU=font_id
            shape.CellsU('Char.Color').FormulaU='RGB(30,41,59)'
            shape.CellsU('Para.HorzAlign').ResultIU=1
            shape.CellsU('VerticalAlign').ResultIU=1
            for cell in ('LeftMargin','RightMargin','TopMargin','BottomMargin'):
                shape.CellsU(cell).FormulaU='0 in'
            if bold:
                shape.CellsU('Char.Style').ResultIU=1
            return shape
        device_shapes={}
        for node in model['nodes']:
            if stencil:
                shape=page.Drop(stencil.Masters.Item(node['master']),node['cx'],node['cy'])
                shape.CellsU('Width').ResultIU=node['w']
                shape.CellsU('Height').ResultIU=node['h']
                shape.CellsU('PinX').ResultIU=node['cx']
                shape.CellsU('PinY').ResultIU=node['cy']
            else:
                shape=page.DrawRectangle(*rect(node))
                shape.CellsU('FillForegnd').FormulaU='RGB(219,234,254)'
                shape.CellsU('LineColor').FormulaU='RGB(37,99,235)'
            shape.Text=''
            shape.NameU=node['id']
            device_shapes[node['name']]=shape
        line_shapes={}
        for link in model['links']:
            line=page.DrawLine(*link['p'],*link['q'])
            line.NameU=link['id']
            for side,cell in [('from','BeginX'),('to','EndX')]:
                line.CellsU(cell).GlueToPos(device_shapes[link[side]],*link[side+'_anchor'])
            line.CellsU('ShapeRouteStyle').FormulaU='GUARD(2)'
            line.CellsU('LinePattern').ResultIU=1
            line.CellsU('LineColor').FormulaU=rgb(link['color'])
            line.CellsU('LineWeight').FormulaU='2.2 pt' if link['category']=='10G' else '2 pt'
            line.CellsU('BeginArrow').ResultIU=0
            line.CellsU('EndArrow').ResultIU=0
            line.SendToBack()
            line_shapes[link['id']]=line
        for tag in model['labels']:
            shape=label(tag['text'],(tag['x']-tag['width']/2,tag['y']-tag['height']/2,
                                    tag['x']+tag['width']/2,tag['y']+tag['height']/2),11.5)
            shape.NameU=tag['id']
            reference=f'Sheet.{line_shapes[tag["link"]].ID}!'
            for axis in ('X','Y'):
                shape.CellsU('Pin'+axis).FormulaU=f'{reference}Begin{axis}+({reference}End{axis}-{reference}Begin{axis})*{tag["fraction"]}'
            shape.CellsU('Angle').FormulaU=reference+'Angle'+('+180 deg' if abs(tag['theta'])>math.pi/2 else '')
            shape.CellsU('LinePattern').ResultIU=1
            shape.CellsU('LineColor').FormulaU=rgb(tag['color'])
            shape.CellsU('LineWeight').FormulaU='.65 pt'
            shape.CellsU('FillPattern').ResultIU=1
            shape.CellsU('FillForegnd').FormulaU='RGB(255,255,255)'
        for node in model['nodes']:
            label(node['name'],node['name_box'],16 if node['role']=='core' else 13,node['role']=='core')
        width,height=model['page_width'],model['page_height']
        label(model['title'],(.8,height-1.3,width-.8,height-.65),28,True)
        label('安全层 · 核心层 · 接入层',(.8,height-1.9,width-.8,height-1.52),14)
        label('线路类型',(1,height-1.05,3.8,height-.71),15,True)
        for index,item in enumerate(model['legend']):
            y=height-1.65-index*.6
            line=page.DrawLine(1,y,1.7,y)
            line.NameU=f'LegendLine_{index}'
            line.CellsU('LineColor').FormulaU=rgb(item['color'])
            line.CellsU('LineWeight').FormulaU='2.2 pt'
            text=label(item['text'],(1.85,y-.16,4.4,y+.16),14)
            text.NameU=f'LegendText_{index}'
            text.CellsU('Para.HorzAlign').ResultIU=0
        for role,title in [('firewall','安全层'),('core','核心层'),('access','接入层')]:
            row=[n for n in model['nodes'] if n['role']==role]
            if row:
                y=row[0]['cy']
                label(title,(.25,y-.2,1.75,y+.15),15,True)
        label('接口标签位于对应线路上；标签边框颜色与线路颜色一致',(1,.85,width-1,1.17),12)
        label('数据来源：'+model['source']['file'],(1,.38,width-1,.7),11)
        background=page.DrawRectangle(0,0,width,height)
        background.CellsU('LinePattern').ResultIU=0
        background.CellsU('FillForegnd').FormulaU='RGB(255,255,255)'
        background.SendToBack()
        doc.SaveAs(str(output))
        if preview:
            page.Export(str(Path(preview).resolve()))
        doc.Close()
        doc=None
        print('Saved:',output,flush=True)
    finally:
        if doc is not None:
            doc.Saved=True
            doc.Close()
        if stencil is not None:
            stencil.Close()
        # Never close a Visio instance containing a user's other documents.
        if app is not None and app.Documents.Count==0:
            app.Quit()
        pythoncom.CoUninitialize()
    return verify_vsdx(output,model)


def main(argv=None):
    parser=argparse.ArgumentParser(description='Excel/CSV to editable Visio network topology')
    parser.add_argument('--version',action='version',version=VERSION)
    commands=parser.add_subparsers(dest='command',required=True)
    masters=commands.add_parser('masters',help='List master names in a .vssx without starting Visio')
    masters.add_argument('stencil')
    planning=commands.add_parser('plan',help='Read links and generate a checked JSON layout')
    planning.add_argument('input')
    planning.add_argument('--output',required=True)
    planning.add_argument('--sheet')
    planning.add_argument('--config')
    planning.add_argument('--stencil')
    rendering=commands.add_parser('render',help='Generate .vsdx and optional PNG using local Visio')
    rendering.add_argument('plan')
    rendering.add_argument('--output',required=True)
    rendering.add_argument('--preview')
    rendering.add_argument('--report')
    rendering.add_argument('--visible',action='store_true')
    rendering.add_argument('--overwrite',action='store_true')
    validation=commands.add_parser('validate',help='Verify a plan and optionally its saved .vsdx')
    validation.add_argument('plan')
    validation.add_argument('--vsdx')
    args=parser.parse_args(argv)
    try:
        if args.command=='masters':
            print(json.dumps(master_names(args.stencil),ensure_ascii=False,indent=2))
        elif args.command=='plan':
            rows,source=read_rows(args.input,args.sheet)
            model=plan(rows,source,read_json(args.config) if args.config else {},args.stencil)
            report=validate_plan(model)
            write_json(args.output,model)
            print(json.dumps(report,ensure_ascii=False,indent=2))
        elif args.command=='render':
            if not args.overwrite and (Path(args.output).exists() or (args.preview and Path(args.preview).exists())):
                raise TopologyError('Output exists; choose another name or pass --overwrite')
            model=read_json(args.plan)
            report=render(model,args.output,args.preview,args.visible)
            if args.report:
                write_json(args.report,report)
            print(json.dumps(report,ensure_ascii=False,indent=2))
        else:
            model=read_json(args.plan)
            report=verify_vsdx(args.vsdx,model) if args.vsdx else validate_plan(model)
            print(json.dumps(report,ensure_ascii=False,indent=2))
    except (TopologyError,OSError,KeyError) as error:
        parser.exit(2,f'Error: {error}\n')


if __name__=='__main__':
    main()
